"""Real file reading: PDF text layer, Word, and Tesseract OCR on a scan and a photo.

The OCR cases are skipped when the tesseract binary is not installed. They use
the 'eng' language when Vietnamese data is absent, so they run anywhere.
"""

import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path

from src.utils.reading.extractors import extract_document_text_with_method

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def a4_scan(lines: list[str]):
    """A4 page at 200 dpi with ~12pt text: what a real scan looks like."""

    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1654, 2339), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(FONT, 33)
    for i, line in enumerate(lines):
        draw.text((150, 200 + i * 60), line, fill="black", font=font)
    return image


class ReadingTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        os.environ["OCR_CACHE_DIR"] = str(cls.tmp / "cache")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    @unittest.skipUnless(Path(FONT).exists(), "needs DejaVu font")
    def test_digital_pdf_is_read_from_its_text_layer(self):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas

        pdfmetrics.registerFont(TTFont("DejaVu", FONT))
        path = self.tmp / "GCN_digital.pdf"
        c = canvas.Canvas(str(path), pagesize=A4)
        c.setFont("DejaVu", 12)
        for i, line in enumerate(["GIẤY CHỨNG NHẬN ĐĂNG KÝ DOANH NGHIỆP",
                                  "Mã số doanh nghiệp: 0101234567",
                                  "Địa chỉ trụ sở chính: Số 1 Láng Hạ, phường Láng Hạ, quận Đống Đa, Hà Nội",
                                  "Vốn điều lệ: 10.000.000.000 đồng",
                                  "Người đại diện theo pháp luật: Trần Thị B — Giám đốc"]):
            c.drawString(60, 780 - i * 22, line)
        c.showPage()
        c.save()
        text, method = extract_document_text_with_method(str(path))
        self.assertEqual(method, "pdf_text_layer")
        self.assertIn("0101234567", text)
        self.assertIn("Đống Đa", text)
        self.assertTrue(text.startswith("--- Page 1 ---"))

    def test_word_paragraphs_and_tables(self):
        import docx

        document = docx.Document()
        document.add_paragraph("THÔNG BÁO TIỀN THUẾ NỢ")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Loại thuế", "Số tiền"
        table.cell(1, 0).text, table.cell(1, 1).text = "GTGT", "125.000.000"
        path = self.tmp / "thong_bao.docx"
        document.save(path)
        text, method = extract_document_text_with_method(str(path))
        self.assertEqual(method, "docx")
        self.assertIn("GTGT\t125.000.000", text)

    @unittest.skipUnless(shutil.which("tesseract") and Path(FONT).exists(), "needs tesseract")
    def test_photo_and_scanned_pdf_go_through_ocr(self):
        langs = subprocess.run(["tesseract", "--list-langs"], capture_output=True, text=True).stdout.split()
        os.environ["OCR_LANG"] = "vie+eng" if "vie" in langs else "eng"
        image = a4_scan(["GIAY CHUNG NHAN DANG KY DOANH NGHIEP", "Ma so doanh nghiep: 0101234567"])
        photo, scan = self.tmp / "gcn_photo.png", self.tmp / "gcn_scan.pdf"
        image.save(photo)
        image.save(scan, "PDF", resolution=200)
        for path, expected in ((photo, "ocr_image"), (scan, "ocr_pdf")):
            with self.subTest(file=path.name):
                text, method = extract_document_text_with_method(str(path))
                self.assertEqual(method, expected)
                self.assertIn("0101234567", text.replace(" ", ""))


if __name__ == "__main__":
    unittest.main()


class WithoutOpenCvTest(unittest.TestCase):
    """A cluster without OpenCV must still import the pipeline and read non-scans."""

    def test_pipeline_imports_and_reads_word_without_cv2(self):
        import subprocess
        import sys

        code = (
            "import sys; sys.modules['cv2'] = None\n"            # make `import cv2` fail
            "from src.agents.supervisor import Supervisor\n"
            "from src.utils.reading.extractors import extract_document_text_with_method as read\n"
            "import docx, tempfile, os\n"
            "d = docx.Document(); d.add_paragraph('THÔNG BÁO TIỀN THUẾ NỢ')\n"
            "p = os.path.join(tempfile.mkdtemp(), 'x.docx'); d.save(p)\n"
            "print(read(p)[1])\n"
            "try:\n"
            "    read(p.replace('x.docx', 'y.png'))\n"
            "except ModuleNotFoundError as e:\n"
            "    print('OCR-ERROR:', e)\n"
        )
        run = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True,
                             env={**os.environ}, cwd=Path(__file__).resolve().parents[1])
        self.assertEqual(run.returncode, 0, run.stderr[-2000:])
        self.assertIn("docx", run.stdout)
        self.assertIn("opencv-python-headless", run.stdout)
