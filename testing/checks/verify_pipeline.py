"""End-to-end run with a fake LLM: folders → classification → extraction → BRD checks → outputs.

No API key and no Tesseract needed (the dossier is .txt files).
"""

import json
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "testing"))

from openpyxl import load_workbook

from corp_income import Settings, run_corp_income
from corp_income.checks import parse_date, within_lookback
from corp_income.documents import catalogue, classify
from fake_llm import FakeLLM

failures = []


def expect(condition, message):
    print(("   ✓ " if condition else "   ✗ ") + message)
    if not condition:
        failures.append(message)


# --- 1. Classification ---------------------------------------------------------
print("1. Phân loại theo thư mục + từ khoá")
types, boxes = catalogue()
expect(len(types) >= 14 and "ho_so_phap_ly" in boxes, f"{len(types)} loại hồ sơ, {len(boxes)} thư mục")
cases = [
    (("GCN_DKDN_ABC.pdf", "", ""), "gcn_dkdn"),
    (("scan01.pdf", "GIẤY CHỨNG NHẬN ĐĂNG KÝ DOANH NGHIỆP", ""), "gcn_dkdn"),
    (("bien_ban_01.pdf", "nội dung bất kỳ", "bao_cao_thuc_dia"), "bao_cao_thuc_dia"),
    (("anh.jpg", "không có từ khoá", "ho_so_phap_ly"), "khac"),
    (("QD_xu_phat_moi_truong.pdf", "", ""), "quyet_dinh_xu_phat"),
    (("tra_cuu_masothue.png", "", ""), "tra_cuu_dkdn"),
]
for (filename, content, box), wanted in cases:
    got = classify(filename, content, box)["type"]
    expect(got == wanted, f"{box or '(gốc)'}/{filename} → {got} (kỳ vọng {wanted})")
# Folder narrows: "dieu le" text inside the tax folder must not become dieu_le.
got = classify("x.pdf", "theo điều lệ công ty, thông báo tiền thuế nợ", "ho_so_thue_bhxh_tranh_chap")["type"]
expect(got == "thong_bao_no_thue", f"thư mục thuế giới hạn ứng viên → {got}")

print("\n2. Ngày tháng, cửa sổ 12 tháng")
as_of = parse_date("13/10/2026")
expect(within_lookback("15/11/2025", as_of) and not within_lookback("01/09/2025", as_of), "12 tháng tính đúng")

# --- 3. Full run ---------------------------------------------------------------
GCN = {"company_name": "CÔNG TY TNHH THƯƠNG MẠI ABC", "tax_code": "0101234567", "legal_form": "tnhh_hai_thanh_vien_tro_len",
       "establishment_date": "12/05/2015", "registration_change_number": 6, "latest_change_date": "20/03/2026",
       "charter_capital": 10, "source_unit": "ty dong", "head_office_address": "Số 1 Láng Hạ, Đống Đa, Hà Nội",
       "legal_representatives": ["Trần Thị B"], "operating_status": "",
       "change_history": [{"date": "20/03/2026", "fields_changed": ["legal_rep", "members"]}]}
LOOKUP = {**GCN, "tax_code": "0101234568", "operating_status": "dang_hoat_dong", "change_history": []}
SITE = {"company_name": "", "tax_code": "", "legal_representatives": [], "source_unit": "dong", "change_history": [],
        "site_visit": {"actual_operation_confirmed": True, "observations": ["Có biển hiệu"]}}
MEMBERS = {"source_unit": "trieu dong", "charter_capital": 10000, "document_date": "20/03/2026",
           "members": [
               {"name": "Nguyễn Văn A", "member_type": "ca_nhan", "capital_amount": 6000, "ownership_pct": 60,
                "acquisition_date": "20/03/2026", "acquisition_method": "nhan_chuyen_nhuong"},
               {"name": "Công ty CP Đầu tư XYZ", "member_type": "phap_nhan", "id_number": "0109999999",
                "capital_amount": 3000, "ownership_pct": 30, "acquisition_date": "12/05/2015",
                "acquisition_method": "gop_von_thanh_lap"}],
           "transfers": [{"date": "20/03/2026", "transferor": "Lê Văn C", "transferee": "Nguyễn Văn A"}]}
REP = {"legal_representatives": [{"name": "Trần Thị B", "title": "Giám đốc", "appointed_date": "20/03/2026"}],
       "previous_representatives": [{"name": "Lê Văn C", "until_date": "20/03/2026"}]}
REP_SITE = {"legal_representatives": [], "actual_operator": {"name": "Nguyễn Văn A", "role": "Chủ tịch HĐTV"},
            "people_met": [{"name": "Nguyễn Văn A", "role": "Chủ tịch HĐTV", "decides_business": True}]}
TAX = {"source_unit": "dong", "tax_debts": [{"authority": "Chi cục Thuế Đống Đa", "amount": 125000000,
                                              "overdue_days": 120, "as_of": "30/09/2026", "status": "con_no"}],
       "social_insurance_debts": [], "administrative_penalties": [], "disputes": []}
SITE_RISK = {"source_unit": "dong", "tax_debts": [], "social_insurance_debts": [], "administrative_penalties": [],
             "disputes": [], "site_risks": ["Bình chữa cháy hết hạn"]}

answers = {
    ("scan_gcn.txt", "legal_basic"): GCN, ("scan_gcn.txt", "shareholders"): MEMBERS, ("scan_gcn.txt", "legal_rep"): REP,
    ("masothue.txt", "legal_basic"): LOOKUP, ("masothue.txt", "legal_rep"): {**REP, "previous_representatives": []},
    ("bien_ban.txt", "legal_basic"): SITE, ("bien_ban.txt", "legal_rep"): REP_SITE,
    ("bien_ban.txt", "liabilities"): SITE_RISK, ("thong_bao_no_thue.txt", "liabilities"): TAX,
}
llm = FakeLLM(answers, wrap_in_fence=True, as_blocks=True)  # the messiest reply shape

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp) / "khdn_abc"
    for rel, text in {
        "ho_so_phap_ly/scan_gcn.txt": "GIẤY CHỨNG NHẬN ĐĂNG KÝ DOANH NGHIỆP CÔNG TY TNHH",
        "ho_so_phap_ly/scan_gcn_copy.txt": "GIẤY CHỨNG NHẬN ĐĂNG KÝ DOANH NGHIỆP CÔNG TY TNHH",
        "tra_cuu_ben_ngoai/masothue.txt": "Thông tin người nộp thuế",
        "bao_cao_thuc_dia/bien_ban.txt": "Đoàn đến làm việc tại trụ sở",
        "ho_so_thue_bhxh_tranh_chap/thong_bao_no_thue.txt": "THÔNG BÁO TIỀN THUẾ NỢ",
        "ho_so_khac/ghi_chu.docm": "x",
    }.items():
        (root / rel).parent.mkdir(parents=True, exist_ok=True)
        (root / rel).write_text(text, encoding="utf-8")

    result = run_corp_income(
        llm, [root], loan_date="13/10/2026", borrower_names=["Nguyễn Văn A"],
        internal_data={"legal_basic": {**GCN, "head_office_address": "Số 5 Cầu Giấy, Hà Nội"}},
        watchlist=[{"name": "Công ty CP Đầu tư XYZ", "tax_code": "0109999999", "list": "WL"}],
        settings=Settings.load(llm_provider="anthropic", llm_model="fake"),
    )
    codes = {n: {f["code"]: f for f in item["findings"]} for n, item in result.items.items()}

    print("\n3. Hồ sơ & thư mục")
    expect(len(result.documents) == 4, f"4 file được xử lý ({len(result.documents)})")
    expect(any("same content" in s for s in result.steps), "file trùng nội dung bị bỏ qua và ghi lại")
    expect(any("unsupported" in s for s in result.steps), "file không hỗ trợ được ghi lại")
    by_name = {d["filename"]: d for d in result.documents}
    expect(by_name["bien_ban.txt"]["type"] == "bao_cao_thuc_dia", "thư mục thực địa → báo cáo thực địa")
    expect(by_name["thong_bao_no_thue.txt"]["passes_done"] == ["liabilities"], "thông báo thuế chỉ chạy mục 4")
    expect(sorted(c for c in llm.calls if c[0] == "scan_gcn.txt") ==
           [("scan_gcn.txt", p) for p in ("legal_basic", "legal_rep", "shareholders")], "GCN chạy mục 1-3")

    print("\n4. Kiểm tra BRD")
    expect(codes[1].get("tax_code_mismatch", {}).get("severity") == "alert", "MST lệch → cảnh báo")
    expect("status_active" in codes[1] and "changes_12m" in codes[1], "trạng thái hoạt động + thay đổi 12 tháng")
    expect("internal_stale" in codes[1], "dữ liệu nội bộ chưa cập nhật")
    expect(codes[2].get("borrower_new_member", {}).get("severity") == "alert", "KH mới nhận vốn góp → cảnh báo")
    expect(codes[2].get("watchlist_member", {}).get("severity") == "alert", "cổ đông thuộc WL")
    expect(result.items[2]["summary"]["members"][0]["capital_amount"] == 6_000_000_000, "6.000 triệu → 6 tỷ đồng")
    expect(codes[3].get("rep_changed_12m", {}).get("severity") == "alert", "đổi người đại diện → cảnh báo")
    expect("operator_differs" in codes[3], "người điều hành thực tế ≠ ĐDPL")
    expect(codes[4].get("thue_debt", {}).get("severity") == "alert", "còn nợ thuế → cảnh báo")
    expect("site_risk" in codes[4], "rủi ro PCCC từ thực địa")

    print("\n5. Kết quả lưu ra")
    paths = result.save(Path(tmp) / "out")
    expect(paths["report"].read_text(encoding="utf-8").startswith("# Corp Income Assessment"), "report.md")
    data = json.loads(paths["json"].read_text(encoding="utf-8"))
    expect(all(d["content"] for d in data["documents"]), "result.json giữ văn bản OCR của từng file")
    ocr_files = sorted(p.name for p in paths["ocr"].iterdir())
    expect(ocr_files == sorted(f"{d['filename']}.txt" for d in result.documents), f"ocr/*.txt: {ocr_files}")
    wb = load_workbook(paths["excel"])
    expect(wb.sheetnames == ["Tóm tắt", "1. Pháp lý cơ bản", "2. Cổ đông", "3. Người đại diện",
                             "4. Nợ thuế - tranh chấp", "Kiểm tra", "Hồ sơ"], f"Excel: {wb.sheetnames}")
    expect(wb["2. Cổ đông"]["D2"].value == 6_000_000_000, "Excel: vốn góp là số (đồng)")
    expect(not result.meta["unknown_placeholders"], "template không có placeholder sai")

print("\n6. Không có LLM → không sập")
bare = run_corp_income(None, texts={"GCN_DKDN.txt": "x"}, loan_date="13/10/2026")
expect(bare.documents[0]["errors"].get("legal_basic") == "No LLM configured.", "lỗi được ghi, không ném exception")

if failures:
    print(f"\n❌ {len(failures)} lỗi")
    sys.exit(1)
print("\n✅ pipeline OK")
