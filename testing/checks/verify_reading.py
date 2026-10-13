"""Real file reading: PDF text layer, Word, and Tesseract OCR on a scanned image/PDF.

Needs reportlab, python-docx, pypdfium2 and the tesseract binary. Uses the
'eng' language if Vietnamese data is not installed, so it runs anywhere.
"""

import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))

from corp_income.reading import read_document

failures = []


def expect(condition, message):
    print(("   ✓ " if condition else "   ✗ ") + message)
    if not condition:
        failures.append(message)


tmp = Path(tempfile.mkdtemp())
os.environ["OCR_CACHE_DIR"] = str(tmp / "cache")

# --- 1. Digital PDF: text layer, no OCR ------------------------------------------
print("1. PDF có text layer")
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

font = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
pdfmetrics.registerFont(TTFont("DejaVu", font))
lines = ["GIẤY CHỨNG NHẬN ĐĂNG KÝ DOANH NGHIỆP", "CÔNG TY TRÁCH NHIỆM HỮU HẠN THƯƠNG MẠI ABC",
         "Mã số doanh nghiệp: 0101234567", "Đăng ký lần đầu: ngày 12 tháng 05 năm 2015",
         "Địa chỉ trụ sở chính: Số 1 Láng Hạ, phường Láng Hạ, quận Đống Đa, Hà Nội",
         "Vốn điều lệ: 10.000.000.000 đồng", "Người đại diện theo pháp luật: Trần Thị B — Giám đốc"]
digital = tmp / "GCN_digital.pdf"
c = canvas.Canvas(str(digital), pagesize=A4)
c.setFont("DejaVu", 12)
for i, line in enumerate(lines):
    c.drawString(60, 780 - i * 22, line)
c.showPage()
c.save()
text, method = read_document(digital)
expect(method == "pdf_text_layer", f"đọc bằng text layer ({method})")
expect("0101234567" in text and "Đống Đa" in text, "giữ đúng MST và dấu tiếng Việt")
expect(text.startswith("--- Page 1 ---"), "có mốc trang cho trường 'page'")

# --- 2. Word ---------------------------------------------------------------------
print("\n2. Word (.docx)")
import docx

document = docx.Document()
document.add_paragraph("THÔNG BÁO TIỀN THUẾ NỢ")
table = document.add_table(rows=2, cols=2)
table.cell(0, 0).text, table.cell(0, 1).text = "Loại thuế", "Số tiền"
table.cell(1, 0).text, table.cell(1, 1).text = "GTGT", "125.000.000"
word = tmp / "thong_bao.docx"
document.save(word)
text, method = read_document(word)
expect(method == "docx" and "GTGT\t125.000.000" in text, "đọc đoạn văn và bảng")

# --- 3. Scanned image + scanned PDF: real Tesseract -------------------------------
print("\n3. Ảnh chụp / PDF scan qua Tesseract")
if not shutil.which("tesseract"):
    print("   (bỏ qua — chưa cài tesseract)")
else:
    langs = subprocess.run(["tesseract", "--list-langs"], capture_output=True, text=True).stdout.split()
    os.environ["OCR_LANG"] = "vie+eng" if "vie" in langs else "eng"
    print(f"   OCR_LANG={os.environ['OCR_LANG']}")
    from PIL import Image, ImageDraw, ImageFont

    # A4 at 200 dpi with ~12pt text: what a real scan looks like. (Giant test
    # lettering is a poor proxy — the ruled-line remover treats long strokes as
    # table borders.)
    image = Image.new("RGB", (1654, 2339), "white")
    draw = ImageDraw.Draw(image)
    body = ImageFont.truetype(font, 33)
    for i, line in enumerate(["GIAY CHUNG NHAN DANG KY DOANH NGHIEP", "Ma so doanh nghiep: 0101234567",
                              "Dia chi tru so chinh: So 1 Lang Ha, Dong Da, Ha Noi"]):
        draw.text((150, 200 + i * 60), line, fill="black", font=body)
    photo = tmp / "gcn_photo.png"
    image.save(photo)
    text, method = read_document(photo)
    expect(method == "ocr_image" and "0101234567" in text.replace(" ", ""), f"ảnh → OCR đọc được MST ({method})")

    scanned = tmp / "gcn_scan.pdf"
    image.save(scanned, "PDF", resolution=200)
    text, method = read_document(scanned)
    expect(method == "ocr_pdf" and "0101234567" in text.replace(" ", ""), f"PDF scan → OCR ({method})")

    os.environ["PDF_TEXT_LAYER"] = "0"
    _, method = read_document(digital)
    expect(method == "ocr_pdf", "PDF_TEXT_LAYER=0 → luôn OCR")
    del os.environ["PDF_TEXT_LAYER"]

shutil.rmtree(tmp, ignore_errors=True)
if failures:
    print(f"\n❌ {len(failures)} lỗi")
    sys.exit(1)
print("\n✅ reading OK")
